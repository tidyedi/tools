# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Michael Schertz

from __future__ import annotations

import io

import openpyxl

from x12_tools.xlsx_export import build_xlsx


def test_header_row_and_data_round_trip() -> None:
    records = [
        {"Ref": "ST01", "Name": "Transaction Set Identifier Code", "Min": 3, "Max": 3},
        {"Ref": "ST02", "Name": "Transaction Set Control Number", "Min": 4, "Max": 9},
    ]
    content = build_xlsx(records)

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Ref", "Name", "Min", "Max")
    assert rows[1] == ("ST01", "Transaction Set Identifier Code", 3, 3)
    assert rows[2] == ("ST02", "Transaction Set Control Number", 4, 9)


def test_header_row_is_bold() -> None:
    content = build_xlsx([{"A": 1}])
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws is not None
    assert ws["A1"].font.bold is True
    assert ws["A2"].font.bold is not True


def test_empty_records_still_produces_a_valid_workbook() -> None:
    content = build_xlsx([])
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws is not None
    assert list(ws.iter_rows(values_only=True)) == []

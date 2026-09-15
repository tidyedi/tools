# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Michael Schertz
"""Turn one table's rows into a real .xlsx workbook -- a bold header row,
frozen so it stays visible while scrolling, and columns sized to their
content, rather than a CSV Excel merely happens to open.

Built server-side (via openpyxl) so the pages never need to load a
client-side spreadsheet library -- matches this project's stance of never
loading a script from outside its own static/ directory.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter

#: Widest a column is allowed to get, in characters -- a long free-text
#: value (a segment name, a raw-segment string) shouldn't blow the sheet out
#: to one unreadable mega-column.
_MAX_COLUMN_WIDTH = 60

#: Matches table-tools.js's HIGHLIGHT_MARKER -- a column's value can wrap a
#: substring in this pair (odd-indexed pieces after splitting) to mark it
#: for bold+color rendering, e.g. the one highlighted token in a Code
#: inventory row's raw segment string.
_HIGHLIGHT_MARKER = "\x01"

#: Rich text in an .xlsx cell can only vary per-character font color, not
#: background fill, so this approximates the on-screen <mark> (amber bg,
#: dark text -- see --highlight-bg/--highlight-text in styles.css) with
#: contrast instead: the matched token stays bold+underlined+full black
#: while everything else in the cell is dimmed to gray, so the eye lands on
#: the token the same way a bright background would, without needing an
#: attention-grabbing color competing against 60-some other cells on screen.
#: rgb is ARGB (opaque alpha "FF" + RRGGBB).
_HIGHLIGHT_FONT = InlineFont(b=True, u="single")
_DIM_FONT = InlineFont(color=Color(rgb="FF9A9A9A"))


def _cell_value(value: Any) -> Any:
    """Turn a marked-up string into a rich-text cell value; anything else
    (including a plain, unmarked string) passes through unchanged."""
    if not isinstance(value, str) or _HIGHLIGHT_MARKER not in value:
        return value
    parts = value.split(_HIGHLIGHT_MARKER)
    blocks: list[TextBlock] = [
        TextBlock(_HIGHLIGHT_FONT if i % 2 else _DIM_FONT, part) for i, part in enumerate(parts) if part
    ]
    return CellRichText(*blocks) if blocks else ""


def _plain_text(value: Any) -> str:
    """The same value's visible text, for column-width sizing -- the
    highlight marker itself takes no visible width."""
    if isinstance(value, str):
        return value.replace(_HIGHLIGHT_MARKER, "")
    return str(value)


def build_xlsx(records: list[dict[str, Any]]) -> bytes:
    """Build a single-sheet workbook from ``records`` (one dict per row,
    every dict sharing the same keys, in column order) and return its bytes.

    An empty ``records`` list still produces a valid, openable workbook --
    just with no header and no rows.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Export"

    if not records:
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    headers = list(records[0].keys())
    ws.append(headers)

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for record in records:
        ws.append([_cell_value(record.get(h)) for h in headers])

    widths = [len(str(h)) for h in headers]
    for record in records:
        for i, h in enumerate(headers):
            widths[i] = max(widths[i], len(_plain_text(record.get(h, ""))))
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, _MAX_COLUMN_WIDTH)

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
